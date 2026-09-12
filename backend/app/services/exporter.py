from datetime import datetime
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.colors import HexColor
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


SECTION_NAMES = {"research": "客户摸底", "requirements": "需求拆解", "capabilities": "能力匹配", "solution": "初步方案", "script": "拜访话术"}

DOCX_BODY_EAST_ASIA_FONT = "宋体"
DOCX_HEADING_EAST_ASIA_FONT = "微软雅黑"
DOCX_LATIN_FONT = "Arial"


def _set_docx_style_font(style, east_asia_font: str, size: float, *, bold: bool = False) -> None:
    """Set explicit Latin and East Asian fonts so Word does not fall back to MS Gothic."""
    style.font.name = DOCX_LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), DOCX_LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), DOCX_LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), east_asia_font)
    r_fonts.set(qn("w:cs"), DOCX_LATIN_FONT)
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        r_fonts.attrib.pop(qn(f"w:{attribute}"), None)


def _configure_docx_typography(doc: Document) -> None:
    _set_docx_style_font(doc.styles["Normal"], DOCX_BODY_EAST_ASIA_FONT, 10.5)
    _set_docx_style_font(doc.styles["List Bullet"], DOCX_BODY_EAST_ASIA_FONT, 10.5)
    _set_docx_style_font(doc.styles["Title"], DOCX_HEADING_EAST_ASIA_FONT, 22, bold=True)
    _set_docx_style_font(doc.styles["Heading 1"], DOCX_HEADING_EAST_ASIA_FONT, 16, bold=True)
    _set_docx_style_font(doc.styles["Heading 2"], DOCX_HEADING_EAST_ASIA_FONT, 12, bold=True)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _prevent_table_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _format_docx_table(table, column_widths) -> None:
    table.autofit = False
    for row in table.rows:
        _prevent_table_row_split(row)
        for cell, width in zip(row.cells, column_widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9EEF2")
        cell._tc.get_or_add_tcPr().append(shading)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.bold = True


def _lines(value, depth=0):
    lines = []
    if isinstance(value, dict):
        for key, item in value.items():
            lines.append((depth, str(key)))
            lines.extend(_lines(item, depth + 1))
    elif isinstance(value, list):
        for item in value:
            if isinstance(item, (dict, list)):
                lines.extend(_lines(item, depth))
            else:
                lines.append((depth, f"• {item}"))
    else:
        lines.append((depth, str(value)))
    return lines


def _add_docx_list(doc: Document, values) -> None:
    for value in values or []:
        doc.add_paragraph(str(value), style="List Bullet")


def _add_docx_labeled_paragraph(doc: Document, label: str, value) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{label}：").bold = True
    paragraph.add_run(str(value if value not in (None, "") else "待补充"))


def _unique_citations(citations):
    unique = []
    seen = set()
    for citation in citations or []:
        identity = (
            citation.get("document_id")
            or citation.get("url")
            or citation.get("filename")
            or citation.get("title")
            or str(citation)
        )
        if identity in seen:
            continue
        seen.add(identity)
        unique.append(citation)
    return unique


def _add_docx_citations(doc: Document, citations) -> None:
    citations = _unique_citations(citations)
    if not citations:
        return
    doc.add_heading("引用来源", level=2)
    for citation in citations:
        source = citation.get("title") or citation.get("filename") or "内部资料"
        location = citation.get("url") or citation.get("filename") or ""
        text = source if not location or location == source else f"{source}  {location}"
        doc.add_paragraph(text)


def _add_docx_requirements(doc: Document, content: dict) -> None:
    for key in ("显性需求", "隐性痛点", "建设期望", "关注事项", "待确认问题", "原文依据"):
        values = content.get(key, [])
        if values:
            doc.add_heading(key, level=2)
            _add_docx_list(doc, values)


def _add_docx_capabilities(doc: Document, content: dict) -> None:
    matches = content.get("匹配结果", [])
    if not matches:
        doc.add_paragraph(content.get("提示") or "当前没有可引用的内部能力依据。")
        return
    for index, match in enumerate(matches, 1):
        doc.add_heading(f"{index}. {match.get('能力', '待确认能力')}", level=2)
        _add_docx_labeled_paragraph(doc, "类别", match.get("类别", "未分类"))
        _add_docx_labeled_paragraph(doc, "匹配理由", match.get("匹配理由"))
        _add_docx_labeled_paragraph(doc, "适用条件", match.get("适用条件"))
        _add_docx_labeled_paragraph(doc, "内部资料原文", match.get("引用"))
    if content.get("提示"):
        doc.add_heading("补充说明", level=2)
        doc.add_paragraph(str(content["提示"]))


def _add_docx_solution(doc: Document, content: dict) -> None:
    if content.get("客户现状"):
        doc.add_heading("客户现状", level=2)
        doc.add_paragraph(str(content["客户现状"]))
    for key in ("建设目标", "方案组合", "建设思路", "预期价值", "风险边界"):
        values = content.get(key, [])
        if values:
            doc.add_heading(key, level=2)
            _add_docx_list(doc, values)


def _add_docx_script(doc: Document, content: dict) -> None:
    settings = content.get("拜访设置", {})
    if settings:
        doc.add_heading("拜访设置", level=2)
        table = doc.add_table(rows=2, cols=3)
        table.style = "Table Grid"
        labels = ("拜访类型", "客户角色", "表达风格")
        keys = ("类型", "客户角色", "表达风格")
        for cell, label in zip(table.rows[0].cells, labels):
            cell.text = label
        for cell, key in zip(table.rows[1].cells, keys):
            cell.text = str(settings.get(key, "待补充"))
        _format_docx_table(table, (Mm(57), Mm(57), Mm(57)))
    for index, stage in enumerate(content.get("阶段", []), 1):
        doc.add_heading(f"{index}. {stage.get('名称', f'阶段 {index}')}", level=2)
        _add_docx_labeled_paragraph(doc, "沟通目标", stage.get("目标"))
        _add_docx_labeled_paragraph(doc, "推荐表达", stage.get("推荐表达"))
        if stage.get("问题"):
            paragraph = doc.add_paragraph()
            paragraph.add_run("建议提问：").bold = True
            _add_docx_list(doc, stage["问题"])
    if content.get("禁止承诺"):
        doc.add_heading("风险与禁止承诺", level=2)
        _add_docx_list(doc, content["禁止承诺"])


def export_docx(path: Path, customer_name: str, artifacts) -> None:
    doc = Document()
    _configure_docx_typography(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = 1
    title.add_run(f"{customer_name}拜访准备材料")
    doc.add_paragraph("本材料由智能体生成并经人工确认。具体能力、报价、工期与服务内容以正式文件为准。")
    for artifact in artifacts:
        doc.add_heading(SECTION_NAMES.get(artifact.type.value, artifact.title), level=1)
        if artifact.type.value == "research":
            if artifact.content.get("客户"):
                doc.add_paragraph(f"客户单位：{artifact.content['客户']}")
            facts = artifact.content.get("结构化档案", [])
            if facts:
                table = doc.add_table(rows=1, cols=3)
                table.style = "Table Grid"
                for cell, value in zip(table.rows[0].cells, ("信息项", "内容", "状态")):
                    cell.text = value
                for fact in facts:
                    cells = table.add_row().cells
                    cells[0].text = str(fact.get("label", "信息项"))
                    cells[1].text = str(fact.get("value", "待补充"))
                    cells[2].text = str(fact.get("status") or fact.get("confidence") or "待核实")
                _format_docx_table(table, (Mm(32), Mm(112), Mm(28)))
                edited = [fact for fact in facts if fact.get("edit_history")]
                if edited:
                    doc.add_heading("人工修改记录", level=2)
                    doc.add_paragraph("以下为原始来源与修改记录，不代表修改后值的当前依据。")
                    for fact in edited:
                        for history in fact.get("edit_history", []):
                            line = f"{fact.get('label', '信息项')}：原值“{history.get('value', '待补充')}”，原状态“{history.get('status', '待确认')}”"
                            doc.add_paragraph(line, style="List Bullet")
                            for source in history.get("sources") or []:
                                doc.add_paragraph(f"原始来源：{source.get('title', '资料')} {source.get('url', '')}")
                            if history.get("source_url") and not history.get("sources"):
                                doc.add_paragraph(f"原始来源：{history['source_url']}")
            for key in ("潜在信息化方向", "待补充"):
                values = artifact.content.get(key, [])
                if values:
                    doc.add_heading(key, level=2)
                    for value in values:
                        doc.add_paragraph(str(value), style="List Bullet")
        elif artifact.type.value == "requirements":
            _add_docx_requirements(doc, artifact.content)
        elif artifact.type.value == "capabilities":
            _add_docx_capabilities(doc, artifact.content)
        elif artifact.type.value == "solution":
            _add_docx_solution(doc, artifact.content)
        elif artifact.type.value == "script":
            _add_docx_script(doc, artifact.content)
        else:
            for depth, line in _lines(artifact.content):
                doc.add_paragraph(line, style="List Bullet" if line.startswith("•") else None)
        _add_docx_citations(doc, artifact.citations)
    doc.save(path)


def export_pdf(path: Path, customer_name: str, artifacts) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    colors = {
        "navy": HexColor("#0B2C49"),
        "cyan": HexColor("#0AA7B3"),
        "cyan_light": HexColor("#E8F7F8"),
        "ink": HexColor("#152C3B"),
        "muted": HexColor("#607480"),
        "line": HexColor("#D8E3E8"),
        "pale": HexColor("#F5F8FA"),
        "orange": HexColor("#E78332"),
        "white": HexColor("#FFFFFF"),
    }
    base = getSampleStyleSheet()
    body = ParagraphStyle("ChineseBody", parent=base["BodyText"], fontName="STSong-Light", fontSize=9.5, textColor=colors["ink"], leading=15, spaceAfter=4)
    small = ParagraphStyle("ChineseSmall", parent=body, fontSize=8, leading=12, textColor=colors["muted"])
    label = ParagraphStyle("ChineseLabel", parent=body, fontSize=8.5, leading=13, textColor=colors["muted"])
    h3 = ParagraphStyle("ChineseH3", parent=body, fontSize=11, leading=16, textColor=colors["navy"], spaceBefore=8, spaceAfter=5)
    title = ParagraphStyle("ChineseTitle", parent=body, fontSize=22, leading=30, textColor=colors["navy"], alignment=TA_CENTER, spaceAfter=7)
    subtitle = ParagraphStyle("ChineseSubtitle", parent=small, alignment=TA_CENTER, spaceAfter=16)
    section_title = ParagraphStyle("SectionTitle", parent=body, fontSize=14, leading=20, textColor=colors["white"])
    section_no = ParagraphStyle("SectionNo", parent=small, fontSize=8, textColor=HexColor("#A7E8EA"))
    bullet = ParagraphStyle("ChineseBullet", parent=body, leftIndent=10, firstLineIndent=-8, bulletIndent=0, spaceAfter=3)

    def p(value, style=body):
        return Paragraph(escape(str(value if value not in (None, "") else "待补充")), style)

    def bullets(items):
        if not isinstance(items, list):
            items = [items]
        return [Paragraph(f"●　{escape(str(item))}", bullet) for item in items]

    def section_header(number, name):
        table = Table(
            [[Paragraph(f"0{number}", section_no), Paragraph(name, section_title)]],
            colWidths=[16*mm, 154*mm],
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors["navy"]),
            ("BACKGROUND", (0, 0), (0, 0), colors["cyan"]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (0, 0), 5*mm),
            ("LEFTPADDING", (1, 0), (1, 0), 5*mm),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4*mm),
            ("TOPPADDING", (0, 0), (-1, -1), 4*mm),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4*mm),
        ]))
        return table

    def key_value_rows(content):
        rows = []
        for key, value in content.items():
            if isinstance(value, list):
                rendered = [p(item, body) for item in value] or [p("待补充", small)]
            elif isinstance(value, dict):
                rendered = [p(f"{subkey}：{subvalue}") for subkey, subvalue in value.items()]
            else:
                rendered = [p(value)]
            rows.append([p(key, label), rendered])
        table = Table(rows, colWidths=[30*mm, 140*mm], repeatRows=0)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors["pale"]),
            ("GRID", (0, 0), (-1, -1), .5, colors["line"]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4*mm),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4*mm),
            ("TOPPADDING", (0, 0), (-1, -1), 3*mm),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3*mm),
        ]))
        return table

    def research_content(content):
        elements = []
        if content.get("客户"):
            elements.extend([p("客户单位", h3), p(content["客户"])])
        facts = content.get("结构化档案", [])
        if facts:
            elements.append(p("结构化档案", h3))
            rows = [[p("信息项", label), p("内容", label), p("可信度", label)]]
            for fact in facts:
                rows.append([p(fact.get("label", "信息项")), p(fact.get("value", "待补充")), p(fact.get("status") or fact.get("confidence", "待核实"), small)])
            table = Table(rows, colWidths=[31*mm, 111*mm, 28*mm], repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors["cyan_light"]),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors["navy"]),
                ("GRID", (0, 0), (-1, -1), .5, colors["line"]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3.5*mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3.5*mm),
                ("TOPPADDING", (0, 0), (-1, -1), 3*mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3*mm),
            ]))
            elements.append(table)
            edited = [fact for fact in facts if fact.get("edit_history")]
            if edited:
                elements.extend([p("人工修改记录", h3), p("以下为原始来源与修改记录，不代表修改后值的当前依据。", small)])
                for fact in edited:
                    for history in fact.get("edit_history", []):
                        elements.append(p(f"●　{fact.get('label', '信息项')}：原值“{history.get('value', '待补充')}”，原状态“{history.get('status', '待确认')}”", bullet))
                        for source in history.get("sources") or []:
                            elements.append(p(f"原始来源：{source.get('title', '资料')} {source.get('url', '')}", small))
                        if history.get("source_url") and not history.get("sources"):
                            elements.append(p(f"原始来源：{history['source_url']}", small))
        for key in ("潜在信息化方向", "待补充"):
            if content.get(key):
                elements.append(p(key, h3))
                elements.extend(bullets(content[key]))
        return elements

    def capability_content(content):
        matches = content.get("匹配结果", [])
        if not matches:
            return [p(content.get("提示", "当前没有可引用的内部能力依据。"))]
        elements = []
        for index, match in enumerate(matches, 1):
            header_table = Table(
                [[p(f"{index}. {match.get('能力', '待确认能力')}", h3), p(match.get("类别", "未分类"), small)]],
                colWidths=[135*mm, 35*mm],
            )
            header_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors["cyan_light"]),
                ("BOX", (0, 0), (-1, 0), .5, colors["line"]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4*mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4*mm),
                ("TOPPADDING", (0, 0), (-1, -1), 3*mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3*mm),
            ]))
            rows = []
            for key in ("匹配理由", "适用条件", "引用"):
                if match.get(key):
                    rows.append([p(key, label), p(match[key])])
            table = Table(rows, colWidths=[30*mm, 140*mm])
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), .5, colors["line"]),
                ("BACKGROUND", (0, 0), (0, -1), colors["pale"]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4*mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4*mm),
                ("TOPPADDING", (0, 0), (-1, -1), 3*mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3*mm),
            ]))
            elements.extend([KeepTogether([header_table, table]), Spacer(1, 4*mm)])
        return elements

    def script_content(content):
        elements = []
        settings_data = content.get("拜访设置", {})
        if settings_data:
            elements.append(key_value_rows(settings_data))
            elements.append(Spacer(1, 4*mm))
        stages = content.get("阶段", [])
        for index, stage in enumerate(stages, 1):
            stage_name = stage.get("名称", f"阶段 {index}")
            rows = [[p(f"{index}", section_no), p(stage_name, h3)]]
            for key in ("目标", "推荐表达", "问题"):
                if stage.get(key):
                    value = stage[key]
                    rendered = [p(item) for item in value] if isinstance(value, list) else p(value)
                    rows.append([p(key, label), rendered])
            table = Table(rows, colWidths=[27*mm, 143*mm])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, 0), colors["cyan"]),
                ("BACKGROUND", (1, 0), (1, 0), colors["cyan_light"]),
                ("GRID", (0, 0), (-1, -1), .5, colors["line"]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4*mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4*mm),
                ("TOPPADDING", (0, 0), (-1, -1), 3*mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3*mm),
            ]))
            elements.extend([KeepTogether([table]), Spacer(1, 3*mm)])
        if content.get("禁止承诺"):
            elements.append(p("风险与禁止承诺", h3))
            elements.extend(bullets(content["禁止承诺"]))
        return elements

    def draw_header_footer(canvas, doc):
        canvas.saveState()
        width, height = A4
        canvas.setStrokeColor(colors["line"])
        canvas.setLineWidth(.5)
        canvas.line(20*mm, height - 13*mm, width - 20*mm, height - 13*mm)
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColor(colors["muted"])
        canvas.drawString(20*mm, height - 10*mm, f"{customer_name}  拜访准备材料")
        canvas.drawRightString(width - 20*mm, 10*mm, f"第 {doc.page} 页")
        canvas.restoreState()

    by_type = {artifact.type.value: artifact for artifact in artifacts}
    story = [
        Spacer(1, 13*mm),
        Paragraph(f"{escape(customer_name)}<br/>拜访准备材料", title),
        Paragraph(f"生成日期：{datetime.now().strftime('%Y年%m月%d日')}　｜　人工确认版本", subtitle),
    ]
    overview = [[p("01", section_no), p("客户摸底"), p("02", section_no), p("需求拆解"), p("03", section_no), p("能力匹配")],
                [p("04", section_no), p("初步方案"), p("05", section_no), p("拜访话术"), p("状态", section_no), p("已确认")]]
    overview_table = Table(overview, colWidths=[13*mm, 43*mm, 13*mm, 43*mm, 13*mm, 45*mm])
    overview_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors["pale"]),
        ("BACKGROUND", (0, 0), (0, -1), colors["cyan"]),
        ("BACKGROUND", (2, 0), (2, -1), colors["cyan"]),
        ("BACKGROUND", (4, 0), (4, -1), colors["cyan"]),
        ("GRID", (0, 0), (-1, -1), .5, colors["line"]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3*mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3*mm),
        ("TOPPADDING", (0, 0), (-1, -1), 3*mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3*mm),
    ]))
    story.extend([overview_table, Spacer(1, 8*mm), p("使用说明", h3), p("本材料由智能体辅助生成并经人工确认。公开信息应结合所列来源复核；具体产品能力、报价、工期与服务内容以正式文件为准。"), Spacer(1, 5*mm)])

    for number, kind in enumerate(("research", "requirements", "capabilities", "solution", "script"), 1):
        artifact = by_type.get(kind)
        if not artifact:
            continue
        story.extend([PageBreak(), section_header(number, SECTION_NAMES[kind]), Spacer(1, 6*mm)])
        if kind == "research":
            story.extend(research_content(artifact.content))
        elif kind == "capabilities":
            story.extend(capability_content(artifact.content))
        elif kind == "script":
            story.extend(script_content(artifact.content))
        else:
            story.append(key_value_rows(artifact.content))
        citations = _unique_citations(artifact.citations)
        if citations:
            story.extend([Spacer(1, 6*mm), p("引用来源", h3)])
            for citation in citations:
                source = citation.get("title") or citation.get("filename") or "内部资料"
                location = citation.get("url") or f"内部知识库 / {citation.get('category', '未分类')}"
                story.append(p(f"{source}　{location}", small))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=20*mm,
        leftMargin=20*mm,
        topMargin=19*mm,
        bottomMargin=17*mm,
        title=f"{customer_name}拜访准备材料",
        author="政企拜访助手",
    )
    doc.build(story, onFirstPage=draw_header_footer, onLaterPages=draw_header_footer)
