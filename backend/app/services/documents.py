from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md"}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)
    if suffix == ".pptx":
        prs = Presentation(path)
        return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
    if suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        return "\n".join(" | ".join(str(v) for v in row if v is not None) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
    raise ValueError("不支持的文件格式")


def chunk_text(text: str, size: int = 700, overlap: int = 100) -> list[str]:
    normalized = " ".join(text.split())
    if not normalized:
        return []
    return [normalized[i:i + size] for i in range(0, len(normalized), max(1, size - overlap))]
