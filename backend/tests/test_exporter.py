import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree

from docx import Document
from pypdf import PdfReader

from app.services.exporter import export_docx, export_pdf


class ExporterTests(unittest.TestCase):
    def setUp(self):
        self.artifact = SimpleNamespace(
            type=SimpleNamespace(value="research"),
            title="客户摸底",
            content={
                "客户": "示例企业",
                "结构化档案": [{
                    "key": "enterprise_nature",
                    "label": "企业性质",
                    "value": "国有企业",
                    "confidence": "公开来源",
                    "status": "已核实",
                    "source_url": "https://example.com",
                    "sources": [{"title": "公开来源", "url": "https://example.com"}],
                    "edit_history": [{
                        "value": "待补充",
                        "status": "待补充",
                        "sources": [{"title": "原始网页", "url": "https://old.example.com"}],
                    }],
                }],
                "潜在信息化方向": ["云网升级"],
                "待补充": ["企业规模"],
            },
            citations=[{"title": "公开来源", "url": "https://example.com"}],
        )

    def test_word_uses_human_readable_research_table_without_internal_keys(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "research.docx"
            export_docx(path, "示例企业", [self.artifact])
            document = Document(path)
            visible = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            research_table = document.tables[0]
            self.assertIn("w:tblHeader", research_table.rows[0]._tr.xml)
            self.assertIn("w:keepNext", research_table.rows[0]._tr.xml)
            for row in research_table.rows:
                self.assertIn("w:cantSplit", row._tr.xml)
        self.assertIn("企业性质", visible)
        self.assertIn("国有企业", visible)
        self.assertIn("已核实", visible)
        self.assertIn("原始来源与修改记录，不代表修改后值的当前依据", visible)
        self.assertIn("原始网页", visible)
        for internal_key in ("enterprise_nature", "source_url", "sources", "label", "value"):
            self.assertNotIn(internal_key, visible)

    def test_word_declares_chinese_fonts_without_ms_gothic_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "research.docx"
            export_docx(path, "示例企业", [self.artifact])
            with zipfile.ZipFile(path) as archive:
                styles_xml = archive.read("word/styles.xml")

        self.assertNotIn(b"MS Gothic", styles_xml)
        root = ElementTree.fromstring(styles_xml)
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        expected = {
            "Normal": ("宋体", "Arial"),
            "Title": ("微软雅黑", "Arial"),
            "Heading1": ("微软雅黑", "Arial"),
            "Heading2": ("微软雅黑", "Arial"),
            "ListBullet": ("宋体", "Arial"),
        }
        for style_id, (east_asia, latin) in expected.items():
            style = root.find(f".//w:style[@w:styleId='{style_id}']", namespace)
            self.assertIsNotNone(style, style_id)
            fonts = style.find("w:rPr/w:rFonts", namespace)
            self.assertIsNotNone(fonts, style_id)
            self.assertEqual(fonts.get(f"{{{namespace['w']}}}eastAsia"), east_asia)
            self.assertEqual(fonts.get(f"{{{namespace['w']}}}ascii"), latin)
            self.assertEqual(fonts.get(f"{{{namespace['w']}}}hAnsi"), latin)
            self.assertIsNone(fonts.get(f"{{{namespace['w']}}}eastAsiaTheme"))

    def test_pdf_accepts_enriched_research_facts(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "research.pdf"
            export_pdf(path, "示例企业", [self.artifact])
            self.assertTrue(path.read_bytes().startswith(b"%PDF"))
            self.assertGreater(path.stat().st_size, 1000)

    def test_exports_hide_internal_capability_fields_and_deduplicate_sources(self):
        document_id = "d85a53ad-043e-40e5-bbe5-cb0552948acd"
        artifact = SimpleNamespace(
            type=SimpleNamespace(value="capabilities"),
            title="能力匹配",
            content={
                "匹配结果": [{
                    "能力": "数据治理咨询",
                    "类别": "产品",
                    "匹配理由": "统一数据口径",
                    "适用条件": "先完成接口摸底",
                    "引用": "提供数据资产目录和数据质量管理能力。",
                    "document_id": document_id,
                }],
                "提示": "能力范围以正式材料为准。",
            },
            citations=[
                {"document_id": document_id, "filename": "数据治理.md", "excerpt": "原文一"},
                {"document_id": document_id, "filename": "数据治理.md", "excerpt": "原文二"},
            ],
        )
        with tempfile.TemporaryDirectory() as directory:
            docx_path = Path(directory) / "capabilities.docx"
            pdf_path = Path(directory) / "capabilities.pdf"
            export_docx(docx_path, "示例企业", [artifact])
            export_pdf(pdf_path, "示例企业", [artifact])
            document = Document(docx_path)
            word_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)

        for visible in (word_text, pdf_text):
            self.assertIn("数据治理咨询", visible)
            self.assertIn("提供数据资产目录和数据质量管理能力", visible)
            self.assertNotIn("document_id", visible)
            self.assertNotIn(document_id, visible)
        self.assertEqual(word_text.count("数据治理.md"), 1)
        self.assertEqual(pdf_text.count("数据治理.md"), 1)
        self.assertIn("1. 数据治理咨询", word_text)
        self.assertIn("内部资料原文：提供数据资产目录和数据质量管理能力", word_text)


if __name__ == "__main__":
    unittest.main()
